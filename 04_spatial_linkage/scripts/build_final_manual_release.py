#!/usr/bin/env python3
"""Build the final manually validated one-sensor-per-WWTW release.

This is deliberately a finalisation layer.  It consumes the latest validated
08c run and never imports, reruns, or edits the automatic matching algorithms.
Manual candidate choice is outcome blind: beta fields are joined only after
the selected sensor identities have been frozen.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import math
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from importlib import metadata
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import geopandas as gpd
import numpy as np
import pandas as pd
import pyogrio
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from shapely.geometry import Point


LOGGER = logging.getLogger("final_wwtw_sensor_manual_release")
BNG = "EPSG:27700"
BRANCH_RELATIVE = Path("project_results/final_wwtw_sensor_manual_release")
# Permit-derived sensor identifiers can contain periods (for example Thames
# IDs such as ``TEMP.1580`` and ``CTCR.1804``).  Keep the token deliberately
# narrow while preserving the complete authoritative sensor_uid.
SENSOR_UID_RE = re.compile(r"[a-z][a-z_]*::[A-Za-z0-9_.-]+")
WWTW_CODE_RE = re.compile(r"UK[A-Z0-9_]+")
SCIENTIFIC_STATEMENT = (
    "Manual inspection was used to identify additional plausible\n"
    "treatment-works-associated CSO sensors where automated name matching failed,\n"
    "including cases where sensor records used road, outlet or local asset names\n"
    "rather than the treatment-works name. The resulting association is a\n"
    "manually validated site-level linkage. It does not by itself establish\n"
    "underground hydraulic connectivity."
)

FINAL_COLUMNS = [
    "company", "uwwCode", "canonical_wwtw_name",
    "wwtw_easting", "wwtw_northing", "wwtw_longitude", "wwtw_latitude",
    "final_assignment_status", "assignment_source",
    "sensor_uid", "permit_number", "sensor_location_name",
    "sensor_easting", "sensor_northing", "sensor_longitude", "sensor_latitude",
    "fitted_beta", "beta_ci_lower_95", "beta_ci_upper_95", "beta_ci_width",
    "eligible_for_beta_analysis", "beta_quality_status", "tail_event_count",
    "load_entering_pe_latest", "design_capacity_pe_latest", "capacity_ratio_latest",
    "capacity_data_year", "sewershed_area_km2",
    "hydro_high_productivity_pct", "hydro_moderate_productivity_pct",
    "hydro_low_productivity_pct", "hydro_no_groundwater_pct",
    "hydro_intergranular_flow_pct", "hydro_fracture_flow_pct",
    "hydro_valid_coverage_pct", "hydro_extraction_status",
]


@dataclass(frozen=True)
class Inputs:
    root: Path
    manual: Path
    branch: Path
    base_run: Path
    base_run_id: str
    base_master: Path
    base_candidates: Path
    sensor_master: Path
    wwtw_master: Path
    place_gpkg: Path


class ReleaseBlocked(RuntimeError):
    """Raised when a release gate prevents publication."""


def parse_args(argv: Iterable[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=".")
    parser.add_argument("--manual-matches", default="data_raw/matching.txt")
    parser.add_argument("--run-id")
    parser.add_argument("--validate-only", action="store_true")
    return parser.parse_args(argv)


def clean(value: object) -> str:
    return "" if pd.isna(value) else str(value).strip()


def valid_xy(easting: object, northing: object) -> bool:
    try:
        return math.isfinite(float(easting)) and math.isfinite(float(northing))
    except (TypeError, ValueError):
        return False


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require_columns(frame: pd.DataFrame, columns: Sequence[str], label: str) -> None:
    missing = sorted(set(columns).difference(frame.columns))
    if missing:
        raise ReleaseBlocked(f"{label} is missing required columns: {missing}")


def read_csv(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise ReleaseBlocked(f"Required CSV does not exist: {path}")
    return pd.read_csv(path, low_memory=False)


def resolve_inputs(root: Path, manual_argument: Path) -> Inputs:
    root = root.resolve()
    manual = manual_argument if manual_argument.is_absolute() else root / manual_argument
    primary_branch = root / "project_results/new_sensor_assignment_approach_primary_spatial"
    latest = primary_branch / "LATEST.txt"
    if not latest.is_file():
        raise ReleaseBlocked(f"Missing validated-primary pointer: {latest}")
    relative = latest.read_text(encoding="utf-8").strip().replace("\\", "/")
    base_run = (primary_branch / relative).resolve()
    try:
        base_run.relative_to((primary_branch / "runs").resolve())
    except ValueError as exc:
        raise ReleaseBlocked("LATEST.txt points outside the primary runs directory") from exc
    manifest_path = base_run / "run_manifest.json"
    if not manifest_path.is_file():
        raise ReleaseBlocked(f"Missing base manifest: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("status") != "validated":
        raise ReleaseBlocked(f"Base run is not validated: {base_run}")
    checksums_path = base_run / "SHA256SUMS.json"
    checksums = json.loads(checksums_path.read_text(encoding="utf-8"))
    for relative_path, expected in checksums.items():
        candidate = base_run / relative_path
        if candidate.name == "SHA256SUMS.json":
            continue
        if not candidate.is_file() or sha256(candidate) != expected:
            raise ReleaseBlocked(f"Base-run checksum failed: {candidate}")
    return Inputs(
        root=root, manual=manual.resolve(), branch=root / BRANCH_RELATIVE,
        base_run=base_run, base_run_id=clean(manifest.get("run_id")) or base_run.name,
        base_master=base_run / "tables/wwtw_primary_sensor_master.csv",
        base_candidates=base_run / "tables/wwtw_primary_sensor_candidates.csv",
        sensor_master=root / "data_intermediate/sensor_master/cso_sensor_master.csv",
        wwtw_master=root / "data_intermediate/catchment_characteristics/catchment_characteristics_master.csv",
        place_gpkg=root / "project_results/new_sensor_assignment_approach/qgis/wwtw_sensor_name_matching.gpkg",
    )


def parse_manual_matches(path: Path) -> tuple[pd.DataFrame, int]:
    if not path.is_file():
        raise ReleaseBlocked(f"Manual evidence does not exist: {path}")
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    records: list[dict[str, object]] = []
    for line_number, raw in enumerate(lines, 1):
        codes = WWTW_CODE_RE.findall(raw)
        sensors = SENSOR_UID_RE.findall(raw)
        data_like = bool(codes or sensors)
        if not data_like:
            continue
        unique_sensors = list(dict.fromkeys(sensors))
        if len(codes) == 1 and unique_sensors:
            status = "parsed_single_sensor" if len(unique_sensors) == 1 else "parsed_multi_candidate"
            notes = "" if len(sensors) == len(unique_sensors) else "duplicate sensor token removed"
            code = codes[0]
        else:
            status = "parse_error"
            notes = f"found {len(codes)} WWTW identifiers and {len(unique_sensors)} sensor identifiers"
            code = codes[0] if len(codes) == 1 else ""
        records.append({
            "source_line_number": line_number, "source_raw_line": raw,
            "uwwCode": code, "listed_sensor_uids": " | ".join(unique_sensors),
            "listed_sensor_count": len(unique_sensors), "parse_status": status,
            "validation_status": "not_validated", "validation_notes": notes,
        })
    columns = [
        "source_line_number", "source_raw_line", "uwwCode", "listed_sensor_uids",
        "listed_sensor_count", "parse_status", "validation_status", "validation_notes",
    ]
    return pd.DataFrame(records, columns=columns), len(lines)


def _manual_sensor_ids(row: Mapping[str, object]) -> list[str]:
    return [part.strip() for part in clean(row["listed_sensor_uids"]).split("|") if part.strip()]


def validate_manual_matches(
    parsed: pd.DataFrame,
    wwtw: pd.DataFrame,
    sensors_identity: pd.DataFrame,
    base: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Validate and select manual IDs without receiving any beta columns."""
    require_columns(wwtw, ["company", "uwwCode"], "authoritative WWTW table")
    require_columns(
        sensors_identity, ["sensor_uid", "company", "bng_easting", "bng_northing"],
        "authoritative sensor identity table",
    )
    require_columns(base, ["company", "uwwCode", "wwtw_easting", "wwtw_northing", "primary_sensor_uid"], "base primary master")
    forbidden = [name for name in sensors_identity.columns if "beta" in name.lower() or name in {"tail_spill_count", "eligible_for_beta_analysis"}]
    if forbidden:
        raise AssertionError(f"Outcome fields supplied to manual candidate selector: {forbidden}")

    result = parsed.copy()
    issues: list[dict[str, object]] = []
    applied: list[dict[str, object]] = []
    w_counts = wwtw["uwwCode"].astype(str).value_counts()
    s_counts = sensors_identity["sensor_uid"].astype(str).value_counts()
    w_index = wwtw.assign(uwwCode=wwtw.uwwCode.astype(str)).set_index("uwwCode", drop=False)
    s_index = sensors_identity.assign(sensor_uid=sensors_identity.sensor_uid.astype(str)).set_index("sensor_uid", drop=False)
    b_index = base.assign(uwwCode=base.uwwCode.astype(str)).set_index("uwwCode", drop=False)
    base_selected = {
        clean(row.primary_sensor_uid): clean(row.uwwCode)
        for row in base.itertuples()
        if clean(row.primary_sensor_uid)
    }
    duplicate_wwtw = set(result.loc[result.uwwCode.astype(str).duplicated(keep=False), "uwwCode"].astype(str))
    listed_flat = [uid for row in result.to_dict("records") for uid in _manual_sensor_ids(row)]
    duplicate_manual_sensors = {uid for uid, count in Counter(listed_flat).items() if count > 1}

    for index, row in result.iterrows():
        row_issues: list[str] = []
        code = clean(row["uwwCode"])
        ids = _manual_sensor_ids(row)
        if row["parse_status"] == "parse_error":
            row_issues.append("manual_parse_error")
        if code in duplicate_wwtw:
            row_issues.append("duplicate_manual_wwtw_record")
        if int(w_counts.get(code, 0)) != 1:
            row_issues.append("manual_wwtw_not_exactly_once")
        if any(uid in duplicate_manual_sensors for uid in ids):
            row_issues.append("duplicate_manual_sensor_record")
        for uid in ids:
            if int(s_counts.get(uid, 0)) != 1:
                row_issues.append(f"manual_sensor_not_exactly_once:{uid}")
        if row_issues:
            status = row_issues[0]
        else:
            w_row = w_index.loc[code]
            w_company = clean(w_row["company"])
            b_row = b_index.loc[code]
            if clean(b_row["company"]) != w_company:
                row_issues.append("base_wwtw_company_inconsistency")
            for uid in ids:
                if clean(s_index.loc[uid, "company"]) != w_company:
                    row_issues.append(f"manual_company_conflict:{uid}")
                assigned_elsewhere = base_selected.get(uid)
                if assigned_elsewhere and assigned_elsewhere != code:
                    row_issues.append(f"manual_sensor_selected_for_other_wwtw:{uid}:{assigned_elsewhere}")
            if not valid_xy(b_row["wwtw_easting"], b_row["wwtw_northing"]):
                row_issues.append("manual_wwtw_geometry_missing")
            if not row_issues:
                candidates = []
                for uid in ids:
                    s_row = s_index.loc[uid]
                    usable = valid_xy(s_row["bng_easting"], s_row["bng_northing"])
                    distance = (
                        math.hypot(float(s_row["bng_easting"]) - float(b_row["wwtw_easting"]),
                                   float(s_row["bng_northing"]) - float(b_row["wwtw_northing"]))
                        if usable else math.nan
                    )
                    candidates.append((uid, usable, distance))
                usable_candidates = [item for item in candidates if item[1]]
                if not usable_candidates:
                    row_issues.append("manual_sensor_geometry_missing")
                else:
                    minimum_distance = min(item[2] for item in usable_candidates)
                    tied = [
                        item for item in usable_candidates
                        if math.isclose(item[2], minimum_distance, rel_tol=1e-12, abs_tol=1e-9)
                    ]
                    selected_uid, _, distance = sorted(tied, key=lambda item: item[0])[0]
                    alternatives = [uid for uid, _, _ in candidates if uid != selected_uid]
                    method = (
                        "explicit_manual_single_sensor" if len(ids) == 1
                        else "nearest_of_manually_approved_candidates_outcome_blind"
                    )
                    applied.append({
                        "uwwCode": code, "wwtw_company": w_company,
                        "wwtw_name": clean(b_row.get("canonical_wwtw_name", w_row.get("uwwName", ""))),
                        "selected_sensor_uid": selected_uid,
                        "selected_sensor_company": clean(s_index.loc[selected_uid, "company"]),
                        "alternative_manual_sensor_uids": " | ".join(alternatives),
                        "previous_automatic_sensor_uid": clean(b_row["primary_sensor_uid"]),
                        "final_assignment_source": "manual_inspection",
                        "selected_sensor_coordinate_available": True,
                        "selection_method": method, "selected_distance_m": distance,
                        "source_line_number": int(row["source_line_number"]),
                        "source_raw_line": row["source_raw_line"],
                    })
            status = row_issues[0] if row_issues else "valid_manual_assignment"
        result.at[index, "validation_status"] = status
        combined = [clean(row["validation_notes"]), *row_issues]
        result.at[index, "validation_notes"] = "; ".join(part for part in combined if part)
        if row_issues:
            issue = row.to_dict()
            issue.update({"validation_status": status, "validation_notes": "; ".join(row_issues)})
            issues.append(issue)

    applied_frame = pd.DataFrame(applied)
    issue_frame = pd.DataFrame(issues, columns=list(result.columns))
    if not applied_frame.empty and applied_frame.selected_sensor_uid.duplicated().any():
        raise ReleaseBlocked("A manually selected sensor was selected for more than one WWTW")
    return result, applied_frame, issue_frame


def freeze_final_sensor_ids(base: pd.DataFrame, applied: pd.DataFrame) -> pd.DataFrame:
    """Freeze assignment identities before any response variable is joined."""
    require_columns(base, ["company", "uwwCode", "primary_sensor_uid"], "base primary master")
    frozen = base[["company", "uwwCode", "primary_sensor_uid"]].copy()
    frozen = frozen.rename(columns={"primary_sensor_uid": "previous_automatic_sensor_uid"})
    frozen["sensor_uid"] = frozen["previous_automatic_sensor_uid"].map(clean).replace("", pd.NA)
    manual_map = {} if applied.empty else applied.set_index("uwwCode")["selected_sensor_uid"].to_dict()
    manual_codes = set(manual_map)
    frozen.loc[frozen.uwwCode.isin(manual_codes), "sensor_uid"] = frozen.loc[frozen.uwwCode.isin(manual_codes), "uwwCode"].map(manual_map)
    frozen["final_assignment_status"] = np.select(
        [frozen.uwwCode.isin(manual_codes), frozen.sensor_uid.notna()],
        ["assigned_manual", "assigned_automatic"],
        default="unmatched_treatment_works",
    )
    frozen["assignment_source"] = np.select(
        [frozen.uwwCode.isin(manual_codes), frozen.sensor_uid.notna()],
        ["manual_inspection", "automatic_primary_spatial"],
        default="none_after_manual_review",
    )
    selected = frozen.sensor_uid.dropna().astype(str)
    if selected.duplicated().any():
        duplicates = sorted(selected[selected.duplicated(keep=False)].unique())
        raise ReleaseBlocked(f"Final selected sensors are not unique: {duplicates[:10]}")
    return frozen


def build_final_tables(
    wwtw: pd.DataFrame,
    base: pd.DataFrame,
    sensors: pd.DataFrame,
    frozen: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Join WWTW predictors and sensor outcomes after identities are frozen."""
    wwtw_fields = [
        "company", "uwwCode", "uwwName", "load_entering_pe_latest", "design_capacity_pe_latest",
        "capacity_ratio_latest", "capacity_data_year", "sewershed_area_km2",
        "hydro_high_productivity_pct", "hydro_moderate_productivity_pct",
        "hydro_low_productivity_pct", "hydro_no_groundwater_pct",
        "hydro_intergranular_flow_pct", "hydro_fracture_flow_pct",
        "hydro_valid_coverage_pct", "hydro_extraction_status",
    ]
    require_columns(wwtw, wwtw_fields, "authoritative WWTW table")
    base_geo = base[["company", "uwwCode", "canonical_wwtw_name", "wwtw_easting", "wwtw_northing", "wwtw_longitude", "wwtw_latitude"]]
    authoritative = wwtw[wwtw_fields].merge(base_geo, on=["company", "uwwCode"], how="left", validate="one_to_one")
    authoritative["canonical_wwtw_name"] = authoritative.canonical_wwtw_name.fillna(authoritative.uwwName)
    authoritative = authoritative.drop(columns="uwwName")
    result = authoritative.merge(
        frozen, on=["company", "uwwCode"], how="left", validate="one_to_one"
    )
    sensor_fields = [
        "sensor_uid", "permit_number", "location_name", "bng_easting", "bng_northing",
        "longitude", "latitude", "fitted_beta", "beta_ci_lower_95", "beta_ci_upper_95",
        "beta_ci_width", "eligible_for_beta_analysis", "beta_quality", "tail_spill_count",
    ]
    require_columns(sensors, sensor_fields, "authoritative sensor master")
    if sensors.sensor_uid.astype(str).duplicated().any():
        raise ReleaseBlocked("Authoritative sensor_uid is not unique")
    sensor_join = sensors[sensor_fields].rename(columns={
        "location_name": "sensor_location_name", "bng_easting": "sensor_easting",
        "bng_northing": "sensor_northing", "longitude": "sensor_longitude",
        "latitude": "sensor_latitude", "beta_quality": "beta_quality_status",
        "tail_spill_count": "tail_event_count",
    })
    result = result.merge(sensor_join, on="sensor_uid", how="left", validate="many_to_one")
    if result.sensor_uid.notna().sum() != result.permit_number.notna().sum():
        raise ReleaseBlocked("A selected sensor did not join exactly to the authoritative sensor master")
    supervisor = result[FINAL_COLUMNS].copy()
    audit = result.copy()
    audit.insert(0, "final_pair_uid", audit.apply(
        lambda row: f"{row['company']}::{row['uwwCode']}::{clean(row['sensor_uid'])}" if clean(row["sensor_uid"]) else "", axis=1
    ))
    return supervisor, audit


def build_sensor_classification(
    sensors: pd.DataFrame,
    base_candidates: pd.DataFrame,
    frozen: pd.DataFrame,
    applied: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    coordinate = sensors.loc[
        [valid_xy(e, n) for e, n in zip(sensors.bng_easting, sensors.bng_northing)]
    ].copy()
    selected_map = frozen.dropna(subset=["sensor_uid"]).set_index("sensor_uid")["uwwCode"].astype(str).to_dict()
    candidate_details: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    for row in base_candidates.itertuples():
        candidate_details[clean(row.sensor_uid)].append((clean(row.uwwCode), clean(row.wwtw_name), "existing_accepted_name_match"))
    if not applied.empty:
        for row in applied.itertuples():
            previous = clean(row.previous_automatic_sensor_uid)
            if previous and previous != clean(row.selected_sensor_uid):
                candidate_details[previous].append((clean(row.uwwCode), clean(row.wwtw_name), "previous_automatic_primary_replaced_manually"))
            for uid in [part.strip() for part in clean(row.alternative_manual_sensor_uids).split("|") if part.strip()]:
                candidate_details[uid].append((clean(row.uwwCode), clean(row.wwtw_name), "alternative_manually_approved_candidate"))
    selected = set(selected_map)
    nonselected = set(candidate_details).difference(selected)
    coordinate_ids = set(coordinate.sensor_uid.astype(str))
    nonselected &= coordinate_ids
    classification = coordinate[["sensor_uid", "company", "permit_number", "location_name", "bng_easting", "bng_northing", "longitude", "latitude"]].copy()
    classification["sensor_classification"] = classification.sensor_uid.map(
        lambda uid: "matching_cso_sensor" if uid in selected else (
            "nonselected_matching_candidate" if uid in nonselected else "unmatched_cso_sensor"
        )
    )
    classification["selected_uwwCode"] = classification.sensor_uid.map(selected_map).fillna("")
    classification["candidate_uwwCodes"] = classification.sensor_uid.map(
        lambda uid: " | ".join(sorted({item[0] for item in candidate_details.get(uid, []) if item[0]}))
    )
    classification["candidate_wwtw_names"] = classification.sensor_uid.map(
        lambda uid: " | ".join(sorted({item[1] for item in candidate_details.get(uid, []) if item[1]}))
    )
    classification["candidate_sources"] = classification.sensor_uid.map(
        lambda uid: " | ".join(sorted({item[2] for item in candidate_details.get(uid, []) if item[2]}))
    )
    nonselected_rows = classification.loc[classification.sensor_classification.eq("nonselected_matching_candidate")].copy()
    nonselected_rows["non_selection_reason"] = "plausible WWTW candidate not selected as final primary sensor"
    return classification, nonselected_rows


def safe_attributes(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.columns:
        if column == "geometry":
            continue
        if pd.api.types.is_bool_dtype(result[column]):
            result[column] = result[column].astype("Int64")
        elif result[column].dtype == object:
            result[column] = result[column].map(lambda value: None if pd.isna(value) else str(value))
    return result


def point_frame(frame: pd.DataFrame, easting: str, northing: str) -> gpd.GeoDataFrame:
    usable = frame.loc[[valid_xy(e, n) for e, n in zip(frame[easting], frame[northing])]].copy()
    geometry = [Point(float(e), float(n)) for e, n in zip(usable[easting], usable[northing])]
    return gpd.GeoDataFrame(safe_attributes(usable), geometry=geometry, crs=BNG)


def write_styles(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    styles = {
        "assigned_treatment_works.qml": ("255,0,0,255", "circle"),
        "matching_cso_sensors.qml": ("0,170,70,255", "circle"),
        "unmatched_cso_sensors.qml": ("255,220,0,255", "circle"),
        "unmatched_treatment_works.qml": ("190,0,190,255", "circle"),
        "nonselected_matching_candidates.qml": ("255,140,0,255", "circle"),
        "wwtw_catchment_areas.qml": ("120,170,230,50", "fill"),
    }
    for name, (color, symbol_type) in styles.items():
        if symbol_type == "circle":
            body = f'<layer class="SimpleMarker"><prop k="name" v="circle"/><prop k="color" v="{color}"/><prop k="outline_color" v="40,40,40,255"/><prop k="size" v="3"/></layer>'
            symbol = f'<symbol name="0" type="marker">{body}</symbol>'
        else:
            body = f'<layer class="SimpleFill"><prop k="color" v="{color}"/><prop k="outline_color" v="70,120,180,255"/><prop k="outline_width" v="0.5"/></layer>'
            symbol = f'<symbol name="0" type="fill">{body}</symbol>'
        text = f'<!DOCTYPE qgis PUBLIC "http://mrcc.com/qgis.dtd" "SYSTEM"><qgis version="3.34"><renderer-v2 type="singleSymbol">{symbol}</renderer-v2></qgis>\n'
        (directory / name).write_text(text, encoding="utf-8")


def write_geopackage(
    path: Path,
    supervisor: pd.DataFrame,
    audit: pd.DataFrame,
    classification: pd.DataFrame,
    nonselected: pd.DataFrame,
    catchments: gpd.GeoDataFrame,
) -> dict[str, int]:
    assigned_rows = audit.loc[audit.final_assignment_status.ne("unmatched_treatment_works")]
    unmatched_rows = audit.loc[audit.final_assignment_status.eq("unmatched_treatment_works")]
    assigned = point_frame(assigned_rows, "wwtw_easting", "wwtw_northing")
    unmatched_wwtw = point_frame(unmatched_rows, "wwtw_easting", "wwtw_northing")
    selected_supervisor = supervisor.loc[supervisor.sensor_uid.notna(), [
        "uwwCode", "canonical_wwtw_name", "sensor_uid", "final_assignment_status",
        "assignment_source", "fitted_beta", "beta_ci_lower_95", "beta_ci_upper_95",
        "beta_quality_status",
    ]]
    matching_data = classification.loc[classification.sensor_classification.eq("matching_cso_sensor")].merge(
        selected_supervisor,
        on="sensor_uid", how="left", validate="one_to_one",
    )
    matching = point_frame(matching_data, "bng_easting", "bng_northing")
    unmatched_sensors = point_frame(
        classification.loc[classification.sensor_classification.eq("unmatched_cso_sensor")],
        "bng_easting", "bng_northing",
    )
    nonselected_points = point_frame(nonselected, "bng_easting", "bng_northing")
    catchment_join = catchments.drop(columns=[column for column in ["assignment_source", "final_assignment_status", "sensor_uid"] if column in catchments.columns]).merge(
        supervisor[["company", "uwwCode", "canonical_wwtw_name", "final_assignment_status", "assignment_source", "sensor_uid"]],
        on=["company", "uwwCode"], how="left", validate="one_to_one",
    )
    catchment_join["catchment_area_km2_calculated"] = catchment_join.geometry.area / 1_000_000.0
    layers = {
        "assigned_treatment_works": assigned,
        "matching_cso_sensors": matching,
        "unmatched_cso_sensors": unmatched_sensors,
        "unmatched_treatment_works": unmatched_wwtw,
        "nonselected_matching_candidates": nonselected_points,
        "wwtw_catchment_areas": gpd.GeoDataFrame(safe_attributes(catchment_join), geometry="geometry", crs=BNG),
    }
    if path.exists():
        path.unlink()
    for name, layer in layers.items():
        layer.to_file(path, layer=name, driver="GPKG", engine="pyogrio")
    return {name: len(layer) for name, layer in layers.items()}


def write_csv(frame: pd.DataFrame, path: Path) -> None:
    frame.to_csv(path, index=False, lineterminator="\n")


def write_xlsx(frame: pd.DataFrame, path: Path) -> None:
    """Write the project-native Python workbook with supervisor-facing styling."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Final WWTW Sensor View"
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append([None if pd.isna(value) else value.item() if isinstance(value, np.generic) else value for value in row])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    widths = {}
    for index, column in enumerate(frame.columns, 1):
        sample = [len(str(column)), *(len(str(value)) for value in frame[column].dropna().head(250))]
        widths[index] = min(max(sample) + 2, 34)
        sheet.column_dimensions[get_column_letter(index)].width = max(widths[index], 12)
    for column in ["wwtw_easting", "wwtw_northing", "sensor_easting", "sensor_northing", "load_entering_pe_latest", "design_capacity_pe_latest", "tail_event_count"]:
        if column in frame.columns:
            for cell in sheet[get_column_letter(frame.columns.get_loc(column) + 1)][1:]:
                cell.number_format = "#,##0.0"
    for column in [name for name in frame.columns if name.endswith("_pct")]:
        for cell in sheet[get_column_letter(frame.columns.get_loc(column) + 1)][1:]:
            cell.number_format = "0.00"
    for column in ["fitted_beta", "beta_ci_lower_95", "beta_ci_upper_95", "beta_ci_width", "capacity_ratio_latest"]:
        if column in frame.columns:
            for cell in sheet[get_column_letter(frame.columns.get_loc(column) + 1)][1:]:
                cell.number_format = "0.0000"
    table = Table(displayName="FinalWWTWSensorView", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    workbook.save(path)


def validate_xlsx_csv_identity(csv_path: Path, xlsx_path: Path) -> None:
    csv_frame = pd.read_csv(csv_path, low_memory=False)
    excel_frame = pd.read_excel(xlsx_path, sheet_name="Final WWTW Sensor View")
    if list(csv_frame.columns) != list(excel_frame.columns) or len(csv_frame) != len(excel_frame):
        raise ReleaseBlocked("CSV/XLSX rows or columns differ")
    # XLSX stores IEEE-754 values while CSV round-trips decimal text; the tiny
    # representation difference is not a data difference.
    pd.testing.assert_frame_equal(csv_frame, excel_frame, check_dtype=False, check_exact=False, rtol=0, atol=1e-9)
    workbook = load_workbook(xlsx_path, read_only=False, data_only=False)
    sheet = workbook["Final WWTW Sensor View"]
    if sheet.freeze_panes != "A2" or sheet.auto_filter.ref is None:
        raise ReleaseBlocked("Workbook usability validation failed")


def load_catchments(inputs: Inputs) -> gpd.GeoDataFrame:
    layers = {name for name, _ in pyogrio.list_layers(inputs.place_gpkg)}
    if "wwtw_sewersheds" not in layers:
        raise ReleaseBlocked("Validated place-core GeoPackage has no wwtw_sewersheds layer")
    catchments = gpd.read_file(inputs.place_gpkg, layer="wwtw_sewersheds", engine="pyogrio")
    if catchments.crs is None or catchments.crs.to_epsg() != 27700:
        raise ReleaseBlocked("Validated WWTW sewersheds are not EPSG:27700")
    require_columns(catchments, ["company", "uwwCode", "geometry"], "validated WWTW sewersheds")
    if catchments.duplicated(["company", "uwwCode"]).any():
        raise ReleaseBlocked("Validated WWTW sewersheds are not unique by company-uwwCode")
    if catchments.geometry.isna().any() or catchments.geometry.is_empty.any() or (~catchments.geometry.is_valid).any():
        raise ReleaseBlocked("Validated WWTW sewersheds contain unusable geometry")
    return catchments


def validate_release(
    supervisor: pd.DataFrame,
    audit: pd.DataFrame,
    classification: pd.DataFrame,
    catchments: gpd.GeoDataFrame,
    gpkg: Path,
    authoritative_count: int,
) -> dict[str, bool]:
    assigned = supervisor.sensor_uid.notna()
    sensor_classes = classification.sensor_classification.value_counts()
    expected_layers = {
        "assigned_treatment_works", "matching_cso_sensors", "unmatched_cso_sensors",
        "unmatched_treatment_works", "nonselected_matching_candidates", "wwtw_catchment_areas",
    }
    actual_layers = {name for name, _ in pyogrio.list_layers(gpkg)}
    gates = {
        "supervisor_row_count": len(supervisor) == authoritative_count,
        "supervisor_key_unique": not supervisor.duplicated(["company", "uwwCode"]).any(),
        "assignment_partition": bool((assigned == supervisor.final_assignment_status.ne("unmatched_treatment_works")).all()),
        "selected_sensor_unique": not supervisor.sensor_uid.dropna().duplicated().any(),
        "unassigned_sensor_fields_blank": not supervisor.loc[~assigned, [name for name in FINAL_COLUMNS if name.startswith("sensor_") or "beta" in name or name == "tail_event_count"]].notna().any().any(),
        "sensor_partition_unique": not classification.sensor_uid.duplicated().any(),
        "sensor_classes_exact": set(sensor_classes.index) <= {"matching_cso_sensor", "nonselected_matching_candidate", "unmatched_cso_sensor"},
        "matching_count_equals_assigned": int(sensor_classes.get("matching_cso_sensor", 0)) == int(assigned.sum()),
        "catchments_unique": not catchments.duplicated(["company", "uwwCode"]).any(),
        "layers_exact": actual_layers == expected_layers,
    }
    for layer_name in expected_layers:
        layer = gpd.read_file(gpkg, layer=layer_name, engine="pyogrio")
        gates[f"{layer_name}_bng"] = layer.crs is not None and layer.crs.to_epsg() == 27700
        gates[f"{layer_name}_geometry"] = not layer.geometry.isna().any() and not layer.geometry.is_empty.any() and bool(layer.geometry.is_valid.all())
    assigned_map = len(gpd.read_file(gpkg, layer="assigned_treatment_works", engine="pyogrio"))
    matching_map = len(gpd.read_file(gpkg, layer="matching_cso_sensors", engine="pyogrio"))
    unmatched_map = len(gpd.read_file(gpkg, layer="unmatched_treatment_works", engine="pyogrio"))
    gates["assigned_equals_matching_map"] = assigned_map == matching_map == int(assigned.sum())
    gates["wwtw_map_partition"] = assigned_map + unmatched_map == authoritative_count
    failures = [name for name, passed in gates.items() if not passed]
    if failures:
        raise ReleaseBlocked(f"Release gates failed: {failures}")
    return gates


def package_versions() -> dict[str, str]:
    result = {"python": platform.python_version()}
    for name in ["pandas", "numpy", "geopandas", "shapely", "pyproj", "pyogrio", "openpyxl"]:
        try:
            result[name] = metadata.version(name)
        except metadata.PackageNotFoundError:
            result[name] = "unavailable"
    return result


def git_provenance(root: Path) -> dict[str, str]:
    try:
        commit = subprocess.run(["git", "rev-parse", "HEAD"], cwd=root, check=True, capture_output=True, text=True).stdout.strip()
        dirty = "dirty" if subprocess.run(["git", "status", "--porcelain"], cwd=root, check=True, capture_output=True, text=True).stdout.strip() else "clean"
        return {"commit": commit, "dirty": dirty}
    except (OSError, subprocess.CalledProcessError):
        return {"commit": "unavailable (not a git work tree)", "dirty": "unavailable"}


def file_inventory(directory: Path, exclude: set[str] | None = None) -> dict[str, str]:
    excluded = exclude or set()
    return {
        str(path.relative_to(directory)).replace("/", "\\"): sha256(path)
        for path in sorted(directory.rglob("*")) if path.is_file() and path.name not in excluded
    }


def write_audit_report(
    path: Path, *, inputs: Inputs, run_id: str, created_at: str,
    source_lines: int, parsed: pd.DataFrame, applied: pd.DataFrame,
    supervisor: pd.DataFrame, classification: pd.DataFrame,
    catchments: gpd.GeoDataFrame, layer_counts: Mapping[str, int],
    source_hashes: Mapping[str, str], output_hashes: Mapping[str, str],
) -> None:
    assigned = supervisor.sensor_uid.notna()
    finite = pd.to_numeric(supervisor.loc[assigned, "fitted_beta"], errors="coerce").map(np.isfinite)
    eligible = supervisor.loc[assigned, "eligible_for_beta_analysis"].fillna(False).astype(bool)
    invalid = parsed.validation_status.ne("valid_manual_assignment")
    classes = classification.sensor_classification.value_counts()
    lines = [
        "FINAL WWTW-SENSOR MANUAL RELEASE AUDIT", "",
        "A. Provenance", f"run ID: {run_id}", f"run timestamp UTC: {created_at}",
        f"git: {json.dumps(git_provenance(inputs.root), sort_keys=True)}",
        f"versions: {json.dumps(package_versions(), sort_keys=True)}",
        f"base run: {inputs.base_run_id} ({inputs.base_run})", f"matching.txt: {inputs.manual}",
        f"source hashes: {json.dumps(source_hashes, sort_keys=True)}",
        f"output hashes before audit/manifest/checksum files: {json.dumps(output_hashes, sort_keys=True)}", "",
        "B. Treatment-work universe", f"total authoritative WWTWs: {len(supervisor)}",
        f"total by company: {supervisor.company.value_counts().sort_index().to_dict()}",
        f"duplicate company-uwwCode keys: {int(supervisor.duplicated(['company','uwwCode']).sum())}",
        f"WWTWs with valid geometry: {sum(valid_xy(e,n) for e,n in zip(supervisor.wwtw_easting,supervisor.wwtw_northing))}",
        f"WWTWs without geometry: {sum(not valid_xy(e,n) for e,n in zip(supervisor.wwtw_easting,supervisor.wwtw_northing))}", "",
        "C. Automatic starting point",
        f"automatically assigned count: {int(supervisor.assignment_source.eq('automatic_primary_spatial').sum()) + int((supervisor.assignment_source.eq('manual_inspection') & supervisor.uwwCode.isin(applied.loc[applied.previous_automatic_sensor_uid.astype(str).str.len().gt(0), 'uwwCode'] if not applied.empty else [])).sum())}",
        f"source run ID and path: {inputs.base_run_id}; {inputs.base_run}", "",
        "D. Manual file", f"source line count: {source_lines}", f"parsed record count: {len(parsed)}",
        f"single-sensor record count: {int(parsed.listed_sensor_count.eq(1).sum())}",
        f"multi-candidate record count: {int(parsed.listed_sensor_count.gt(1).sum())}",
        f"invalid record count: {int(invalid.sum())}",
        f"validation statuses: {parsed.validation_status.value_counts().to_dict()}", "",
        "E. Manual changes", f"manual assignments successfully applied: {len(applied)}",
        f"automatic assignments replaced: {int(applied.previous_automatic_sensor_uid.astype(str).str.len().gt(0).sum()) if not applied.empty else 0}",
        f"alternative manual candidates retained: {sum(bool(clean(v)) for v in applied.alternative_manual_sensor_uids) if not applied.empty else 0}",
        f"assignments by company: {applied.wwtw_company.value_counts().sort_index().to_dict() if not applied.empty else {}}",
        "complete applied manual assignments:",
        *(f"  {row.uwwCode} -> {row.selected_sensor_uid}" for row in applied.itertuples()), "",
        "F. Final assignment", f"assigned WWTWs: {int(assigned.sum())} ({assigned.mean()*100:.2f}%)",
        f"unmatched WWTWs: {int((~assigned).sum())} ({(~assigned).mean()*100:.2f}%)",
        f"assigned automatically: {int(supervisor.assignment_source.eq('automatic_primary_spatial').sum())}",
        f"assigned manually: {int(supervisor.assignment_source.eq('manual_inspection').sum())}",
        f"total selected sensors: {supervisor.sensor_uid.notna().sum()}",
        f"assigned WWTW count equals selected sensor count: {int(assigned.sum()) == supervisor.sensor_uid.notna().sum()}", "",
        "G. Beta coverage (calculated after identity selection was frozen)",
        f"selected sensors with finite fitted_beta: {int(finite.sum())}/{int(assigned.sum())} ({finite.mean()*100 if len(finite) else 0:.2f}%)",
        f"selected sensors eligible_for_beta_analysis: {int(eligible.sum())}/{int(assigned.sum())} ({eligible.mean()*100 if len(eligible) else 0:.2f}%)",
        f"all WWTWs with assigned finite-beta sensor: {int(finite.sum())}/{len(supervisor)} ({finite.sum()/len(supervisor)*100:.2f}%)", "",
        "H. Final sensor classes", f"counts: {classes.to_dict()}", f"coordinate-bearing sensor universe: {len(classification)}",
        f"overlap checks: duplicate sensor_uid count={int(classification.sensor_uid.duplicated().sum())}", "",
        "I. Catchments", f"source: {inputs.place_gpkg}; layer=wwtw_sewersheds (Hoffmann-derived)",
        f"source/output CRS: EPSG:27700", f"raw/dissolved WWTW catchment count: {len(catchments)}",
        f"invalid geometry count: {int((~catchments.geometry.is_valid).sum())}",
        f"area km2 summary: {(catchments.geometry.area/1_000_000).describe().to_dict()}", "",
        "J. GeoPackage", f"layer counts: {dict(layer_counts)}", "all layers: EPSG:27700; null geometries: 0",
        f"assigned-treatment-works versus matching-sensors equality: {layer_counts['assigned_treatment_works'] == layer_counts['matching_cso_sensors']}", "",
        "K. Remaining unmatched treatment works",
        *(f"  {row.company}\t{row.uwwCode}\t{row.canonical_wwtw_name}\tno credible sensor after automatic and manual review" for row in supervisor.loc[~assigned].itertuples()), "",
        "L. Scientific statement", SCIENTIFIC_STATEMENT,
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def branch_readme() -> str:
    return """# Final WWTW–CSO sensor manual release

This immutable-run branch finalises the validated automatic primary spatial assignment by applying explicit evidence from `data_raw/matching.txt`. The automatic name matcher and spatial selector are inputs and are neither rerun nor modified.

Precedence is valid manual inspection, then the existing automatic primary assignment, then no assignment. Multi-candidate manual rows are resolved using coordinate availability and nearest EPSG:27700 distance, with lexical `sensor_uid` only for an exact technical tie. Beta and beta availability are prohibited from selection and are joined only after sensor identities are frozen.

The supervisor CSV/XLSX has exactly one row per authoritative company–`uwwCode`, concise WWTW predictors, final assignment fields, and the selected individual sensor response. Detailed parsing, validation, replacement, and sensor-class evidence remains in separate audit tables.

`qgis/final_wwtw_sensor_assignment.gpkg` contains six EPSG:27700 layers: assigned treatment works, matching sensors, unmatched sensors, unmatched treatment works, nonselected plausible candidates, and validated WWTW catchment areas. Catchments reuse the repository's validated Hoffmann-derived `wwtw_sewersheds` layer; no buffers or fabricated polygons are used. A manual linkage is site-level evidence and does not prove underground hydraulic connectivity.

Reproduce with:

```powershell
python scripts/08e_build_final_manual_wwtw_sensor_release.py --project-root . --manual-matches data_raw/matching.txt --validate-only
python scripts/08e_build_final_manual_wwtw_sensor_release.py --project-root . --manual-matches data_raw/matching.txt
```

Validated releases are under `runs/<run_id>/`; `LATEST.txt` changes only after every release gate passes. A run includes tables, QGIS data/styles, diagnostics, logs, the byte-identical manual snapshot, a manifest, and SHA-256 inventory. Invalid manual evidence blocks publication. Remaining limitations are incomplete credible sensor coverage and the interpretive limit of site-level linkage.
"""


def preflight(inputs: Inputs) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, int, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    parsed, source_lines = parse_manual_matches(inputs.manual)
    wwtw = read_csv(inputs.wwtw_master)
    sensors = read_csv(inputs.sensor_master)
    base = read_csv(inputs.base_master)
    identity_columns = ["sensor_uid", "company", "bng_easting", "bng_northing"]
    validated, applied, issues = validate_manual_matches(parsed, wwtw, sensors[identity_columns], base)
    return validated, applied, issues, source_lines, wwtw, sensors, base


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    try:
        inputs = resolve_inputs(Path(args.project_root), Path(args.manual_matches))
        parsed, applied, issues, source_lines, wwtw, sensors, base = preflight(inputs)
        LOGGER.info("Manual records: parsed=%d valid=%d invalid=%d", len(parsed), len(applied), len(issues))
        if not issues.empty:
            for row in issues.itertuples():
                LOGGER.error("line %s %s: %s", row.source_line_number, row.uwwCode, row.validation_notes)
            raise ReleaseBlocked(f"Preflight found {len(issues)} blocking manual record(s)")
        if args.validate_only:
            LOGGER.info("Preflight passed; no release files were written")
            return 0

        run_id = args.run_id or datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ_manual_final")
        if not re.fullmatch(r"[A-Za-z0-9_.-]+", run_id):
            raise ReleaseBlocked("--run-id may contain only letters, digits, period, underscore, and hyphen")
        final_run = inputs.branch / "runs" / run_id
        if final_run.exists():
            raise ReleaseBlocked(f"Run already exists and will not be overwritten: {final_run}")
        inputs.branch.mkdir(parents=True, exist_ok=True)
        (inputs.branch / "runs").mkdir(exist_ok=True)
        staging = Path(tempfile.mkdtemp(prefix=f".{run_id}.staging-", dir=inputs.branch / "runs"))
        try:
            for relative in ["tables", "qgis/styles", "diagnostics", "logs", "source_snapshots"]:
                (staging / relative).mkdir(parents=True, exist_ok=True)
            shutil.copyfile(inputs.manual, staging / "source_snapshots/matching.txt")
            if (staging / "source_snapshots/matching.txt").read_bytes() != inputs.manual.read_bytes():
                raise ReleaseBlocked("matching.txt snapshot is not byte-identical")

            frozen = freeze_final_sensor_ids(base, applied)
            supervisor, audit = build_final_tables(wwtw, base, sensors, frozen)
            base_candidates = read_csv(inputs.base_candidates)
            classification, nonselected = build_sensor_classification(sensors, base_candidates, frozen, applied)
            catchments = load_catchments(inputs)

            tables = staging / "tables"
            write_csv(supervisor, tables / "final_wwtw_sensor_supervisor_view.csv")
            write_xlsx(supervisor, tables / "final_wwtw_sensor_supervisor_view.xlsx")
            write_csv(parsed, tables / "manual_matches_parsed.csv")
            write_csv(applied, tables / "manual_matches_applied.csv")
            write_csv(issues, tables / "manual_match_issues.csv")
            write_csv(audit, tables / "final_wwtw_sensor_assignment_audit.csv")
            write_csv(classification, tables / "final_sensor_layer_classification.csv")
            validate_xlsx_csv_identity(tables / "final_wwtw_sensor_supervisor_view.csv", tables / "final_wwtw_sensor_supervisor_view.xlsx")

            gpkg = staging / "qgis/final_wwtw_sensor_assignment.gpkg"
            layer_counts = write_geopackage(gpkg, supervisor, audit, classification, nonselected, catchments)
            write_styles(staging / "qgis/styles")
            gates = validate_release(supervisor, audit, classification, catchments, gpkg, len(wwtw))

            source_paths = [
                inputs.manual, inputs.base_master, inputs.base_candidates, inputs.sensor_master,
                inputs.wwtw_master, inputs.place_gpkg,
                inputs.root / "scripts/08_match_treatment_works_sensors_by_name.py",
                inputs.root / "scripts/08b_rebuild_wwtw_name_matching_place_core.py",
                inputs.root / "scripts/08c_select_primary_wwtw_sensor_by_distance.py",
                inputs.root / "config/wwtw_sensor_place_core_matching.json",
                inputs.root / "config/wwtw_primary_sensor_selection.json",
            ]
            source_hashes = {str(path.relative_to(inputs.root)).replace("/", "\\"): sha256(path) for path in source_paths}
            created_at = datetime.now(timezone.utc).isoformat()
            preliminary_hashes = file_inventory(staging, {"run_manifest.json", "SHA256SUMS.json"})
            write_audit_report(
                staging / "diagnostics/final_wwtw_sensor_manual_release_audit.txt",
                inputs=inputs, run_id=run_id, created_at=created_at, source_lines=source_lines,
                parsed=parsed, applied=applied, supervisor=supervisor, classification=classification,
                catchments=catchments, layer_counts=layer_counts, source_hashes=source_hashes,
                output_hashes=preliminary_hashes,
            )
            log_text = f"validated run {run_id}\nmanual records {len(parsed)}\nmanual applied {len(applied)}\n"
            (staging / "logs/final_wwtw_sensor_manual_release.log").write_text(log_text, encoding="utf-8")
            manifest = {
                "run_id": run_id, "created_at_utc": created_at, "status": "validated",
                "source_script": "scripts\\08e_build_final_manual_wwtw_sensor_release.py",
                "base_run_id": inputs.base_run_id, "base_run_path": str(inputs.base_run),
                "source_hashes": source_hashes, "matching_txt_sha256": sha256(inputs.manual),
                "metrics": {
                    "authoritative_wwtw_count": len(supervisor), "manual_records": len(parsed),
                    "manual_applied": len(applied), "assigned_wwtw_count": int(supervisor.sensor_uid.notna().sum()),
                    "unmatched_wwtw_count": int(supervisor.sensor_uid.isna().sum()),
                    "coordinate_sensor_count": len(classification), "layer_counts": layer_counts,
                },
                "release_gates": gates, "git": git_provenance(inputs.root), "packages": package_versions(),
            }
            (staging / "run_manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
            checksums = file_inventory(staging, {"SHA256SUMS.json"})
            (staging / "SHA256SUMS.json").write_text(json.dumps(checksums, indent=2, sort_keys=True) + "\n", encoding="utf-8")
            os.replace(staging, final_run)
            (inputs.branch / "README.md").write_text(branch_readme(), encoding="utf-8")
            latest_temp = inputs.branch / ".LATEST.tmp"
            latest_temp.write_text(f"runs/{run_id}\n", encoding="utf-8")
            os.replace(latest_temp, inputs.branch / "LATEST.txt")
            LOGGER.info("Published validated release: %s", final_run)
            return 0
        finally:
            if staging.exists():
                shutil.rmtree(staging)
    except (ReleaseBlocked, AssertionError, pd.errors.MergeError) as exc:
        LOGGER.error("RELEASE BLOCKED: %s", exc)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
